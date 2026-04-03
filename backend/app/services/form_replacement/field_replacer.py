"""
Field Replacer - Thay thế dot-lines bằng input fields HTML
"""

import re
from typing import List, Dict, Tuple
from docx import Document
from docx.shared import Pt, RGBColor
from .dot_line_detector import DotLinePlaceholder, DotLineDetector


class HTMLFieldReplacer:
    """Thay thế dot-lines bằng input fields HTML"""

    @staticmethod
    def _get_input_type(field_type: str) -> str:
        """Convert field_type to HTML input type"""
        type_map = {
            'date': 'date',
            'number': 'number',
            'phone': 'tel',
            'email': 'email',
            'text': 'text',
        }
        return type_map.get(field_type, 'text')

    @staticmethod
    def replace_dot_line_in_html(
        text: str,
        placeholder: DotLinePlaceholder,
        field_id: str,
        field_type: str
    ) -> str:
        """Replace một dot-line bằng HTML input field
        
        Example:
            Input: "Họ tên: ..............................."
            Output: "Họ tên: <input type='text' id='field_name' ...>"
        """
        input_html = (
            f'<input type="{HTMLFieldReplacer._get_input_type(field_type)}" '
            f'id="{field_id}" '
            f'name="{placeholder.field_name}" '
            f'class="dot-line-input dot-line-input-{field_type}" '
            f'placeholder="Nhập {placeholder.label.lower()}..." '
            f'data-field-name="{placeholder.field_name}" '
            f'onFocus="loadSuggestions(event)" '
            f'onInput="filterSuggestions(event)" '
            f'onBlur="hideSuggestionsOnBlur(event)" />'
        )
        
        # Replace dot-line pattern with input
        dot_pattern = text[placeholder.dot_start_pos:placeholder.dot_end_pos]
        result = text[:placeholder.dot_start_pos] + input_html + text[placeholder.dot_end_pos:]
        
        return result

    @staticmethod
    def replace_in_document_html(
        original_text: str,
        placeholders: List[DotLinePlaceholder],
        fields_info: List[Dict]
    ) -> str:
        """Replace tất cả dot-lines trong original text bằng HTML inputs
        
        Giữ nguyên cấu trúc document, chỉ thay thế dot-lines
        """
        # Build map từ placeholder label → field info
        field_map = {f['placeholder']['label']: f for f in fields_info}
        
        # Sort placeholders by position (reverse để không lộn xộn vị trí)
        sorted_placeholders = sorted(placeholders, key=lambda p: p.dot_start_pos, reverse=True)
        
        result = original_text
        
        for placeholder in sorted_placeholders:
            if placeholder.label in field_map:
                field_info = field_map[placeholder.label]
                field_id = f"field_{field_info['name']}"
                
                # Replace dot-line in current line
                before_text = result[:placeholder.dot_start_pos]
                dot_pattern = result[placeholder.dot_start_pos:placeholder.dot_end_pos]
                after_text = result[placeholder.dot_end_pos:]
                
                input_html = (
                    f'<input type="{HTMLFieldReplacer._get_input_type(field_info["field_type"])}" '
                    f'id="{field_id}" '
                    f'name="{field_info["name"]}" '
                    f'class="dot-line-input dot-line-input-{field_info["field_type"]}" '
                    f'placeholder="Nhập {field_info["label"].lower()}..." '
                    f'data-field-name="{field_info["name"]}" />'
                )
                
                result = before_text + input_html + after_text
        
        return result

    @staticmethod
    def render_form_with_replacements(
        paragraphs: List[str],
        placeholders: List[DotLinePlaceholder],
        fields_info: List[Dict]
    ) -> str:
        """Render form HTML bằng cách thay thế dot-lines
        
        Giữ nguyên cấu trúc paragraph, chỉ replace dot-lines
        """
        html_lines = []
        
        # Build map từ line_index → placeholders
        placeholders_by_line = {}
        for p in placeholders:
            if p.line_index not in placeholders_by_line:
                placeholders_by_line[p.line_index] = []
            placeholders_by_line[p.line_index].append(p)
        
        # Build map từ label → field info
        field_map = {f['placeholder']['label']: f for f in fields_info}
        
        # Process each paragraph
        for line_idx, para_text in enumerate(paragraphs):
            if line_idx in placeholders_by_line:
                # This line has dot-lines
                result_text = para_text
                
                # Sort by position (reverse) để không lộn vị trí
                line_placeholders = sorted(
                    placeholders_by_line[line_idx],
                    key=lambda p: p.dot_start_pos,
                    reverse=True
                )
                
                for placeholder in line_placeholders:
                    if placeholder.label in field_map:
                        field_info = field_map[placeholder.label]
                        field_id = f"field_{field_info['name']}"
                        
                        # Create input HTML
                        input_html = (
                            f'<input type="{HTMLFieldReplacer._get_input_type(field_info["field_type"])}" '
                            f'id="{field_id}" '
                            f'name="{field_info["name"]}" '
                            f'class="dot-line-input dot-line-input-{field_info["field_type"]}" '
                            f'placeholder="Nhập {field_info["label"].lower()}..." '
                            f'data-field-name="{field_info["name"]}" '
                            f'onFocus="loadSuggestions(event)" '
                            f'onInput="filterSuggestions(event)" '
                            f'onBlur="hideSuggestionsOnBlur(event)" />'
                        )
                        
                        # Replace dot-line
                        before = result_text[:placeholder.dot_start_pos]
                        after = result_text[placeholder.dot_end_pos:]
                        result_text = before + input_html + after
                
                # Wrap in paragraph
                html_lines.append(f'<p class="form-paragraph">{result_text}</p>')
            else:
                # Normal paragraph
                if para_text.strip():
                    html_lines.append(f'<p class="form-paragraph">{para_text}</p>')
        
        # Wrap in form
        form_html = f'''
        <form id="dot-line-form" class="dot-line-form">
            <div class="form-content">
                {''.join(html_lines)}
            </div>
            <div class="form-actions">
                <button type="submit" class="btn btn-primary">Submit</button>
                <button type="reset" class="btn btn-secondary">Reset</button>
            </div>
        </form>
        '''
        
        return form_html


class DocxFieldReplacer:
    """Thay thế dot-lines trong DOCX document"""

    @staticmethod
    def replace_in_docx(
        input_file: str,
        output_file: str,
        placeholders: List[DotLinePlaceholder]
    ) -> None:
        """Replace dot-lines trong DOCX file bằng text markers
        
        Note: DOCX files không có input elements, chỉ replace bằng brackets
        """
        doc = Document(input_file)
        
        # Build map từ line_index → placeholders
        placeholders_by_line = {}
        for p in placeholders:
            if p.line_index not in placeholders_by_line:
                placeholders_by_line[p.line_index] = []
            placeholders_by_line[p.line_index].append(p)
        
        # Process paragraphs
        for para_idx, para in enumerate(doc.paragraphs):
            if para_idx in placeholders_by_line:
                text = para.text
                
                # Sort by position (reverse)
                line_placeholders = sorted(
                    placeholders_by_line[para_idx],
                    key=lambda p: p.dot_start_pos,
                    reverse=True
                )
                
                for placeholder in line_placeholders:
                    # Replace dot-line with brackets
                    before = text[:placeholder.dot_start_pos]
                    after = text[placeholder.dot_end_pos:]
                    
                    # Create marker for this field
                    marker = f'[{placeholder.label}]'
                    text = before + marker + after
                
                # Update paragraph
                para.text = text
        
        # Save
        doc.save(output_file)
